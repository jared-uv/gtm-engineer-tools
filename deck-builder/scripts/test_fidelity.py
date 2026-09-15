#!/usr/bin/env python3
"""Tests for the parts of the fidelity check and the exporter that don't need Chrome:
counting slides, the preview's CSS, font mapping, the image comparison on synthetic
renders, explaining regions, verdicts, and how an export decides a slide's treatment.

    python3 test_fidelity.py
Needs Pillow for the comparison tests; skips them without it.
"""
import json, sys, tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import fidelity, stage

passed = failed = 0
def check(name, cond, detail=""):
    global passed, failed
    if cond: passed += 1; print(f"  ok   {name}")
    else: failed += 1; print(f"  FAIL {name}  {detail}")

# ── slide counting ──────────────────────────────────────────────────────────
print("slide_count")
check("slide-deck children", fidelity.slide_count("<slide-deck><section>a</section><section>b</section></slide-deck>") == 2)
check("deck-stage children", fidelity.slide_count("<deck-stage>\n<section data-label='x'>a</section>\n<section>b</section><section>c</section></deck-stage>") == 3)
check("nested sections don't count", fidelity.slide_count("<slide-deck><section><section>in</section></section><section>b</section></slide-deck>") == 2)
check("no deck element falls back to all top-level sections", fidelity.slide_count("<body><section>a</section></body>") == 1)

# ── preview CSS ─────────────────────────────────────────────────────────────
print("preview css")
c = {"r": 10, "g": 13, "b": 27, "a": 1}
check("color", fidelity.css_color(c) == "rgba(10,13,27,1)")
check("missing color is transparent", fidelity.css_color(None) == "transparent")
g = {"kind": "linear", "angle": 90, "stops": [{"c": c, "p": 0}, {"c": {"r": 255, "g": 0, "b": 0, "a": 1}, "p": 1}]}
check("linear gradient", fidelity.css_gradient(g).startswith("linear-gradient(90deg, rgba(10,13,27,1) 0.0%"))
check("radial gradient, no centre measured", fidelity.css_gradient({**g, "kind": "radial"}).startswith("radial-gradient(circle at 50.0% 50.0%,"))
check("radial gradient keeps its centre and shape", fidelity.css_gradient({**g, "kind": "radial", "shape": "ellipse", "at": {"x": 0.5, "y": 0.6}}).startswith("radial-gradient(ellipse at 50.0% 60.0%,"))
three = {**g, "stops": g["stops"][:1] + [{"c": {"r": 100, "g": 44, "b": 169, "a": 1}, "p": 0.7}] + g["stops"][1:]}
check("every stop reaches the preview", "rgba(100,44,169,1) 70.0%" in fidelity.css_gradient(three))
tm = {"items": [{"type": "text", "x": 0, "y": 0, "w": 600, "h": 297, "align": "left", "valign": "top", "lineHeight": 98.8, "wrap": True,
                 "runs": [{"t": "FOUR", "family": "Inter", "size": 104, "weight": 900, "italic": True, "color": c, "spacing": 0}]}], "background": None}
check("a text box's strut uses its own font size, not the browser's 16px",
      "font-size:104px;font-family:'Inter'" in fidelity.preview_html(tm, {}, "", "/").split("<span")[0])

# ── font mapping ────────────────────────────────────────────────────────────
print("fonts")
brand = {"fonts": {"mono": {"family": "SF Mono", "slides_equivalent": "Roboto Mono"},
                   "display": {"family": "Acme Grotesk", "slides_equivalent": "Archivo"}}}
check("brand equivalent wins", fidelity.font_for("SF Mono", brand) == "Roboto Mono")
check("brand equivalent, any case", fidelity.font_for("acme grotesk", brand) == "Archivo")
check("a Slides font stays", fidelity.font_for("Inter", {}) == "Inter")
check("unknown mono becomes Roboto Mono", fidelity.font_for("Menlo", {}) == "Roboto Mono")
check("unknown sans becomes Arial", fidelity.font_for("Proxima Nova", {}) == "Arial")
check("generic for mono", fidelity.generic_for("Roboto Mono") == "monospace")
check("generic for sans", fidelity.generic_for("Inter") == "sans-serif")
m = {"items": [{"type": "text", "x": 10, "y": 10, "w": 100, "h": 20,
                "runs": [{"t": "ls", "family": "SF Mono"}, {"t": "\n", "br": True}, {"t": "x", "family": "Inter"}]}]}
subs = fidelity.substitutions(m, brand)
check("substitution found for SF Mono only", len(subs) == 1 and subs[0]["effects"] == ["font not in Google Slides: SF Mono → Roboto Mono"], subs)
tracked = {"items": [{"type": "text", "x": 0, "y": 0, "w": 10, "h": 10, "runs": [{"t": "TITLE", "family": "Inter", "spacing": -1.04, "weight": 900}]}]}
tr = fidelity.substitutions(tracked, {})
check("letter spacing and black weight are named for Google Slides",
      tr and tr[0]["effects"] == ["letter spacing not in Google Slides", "weight 900 drawn as bold"], tr)
check("letter spacing is kept for PowerPoint", fidelity.substitutions(tracked, {}, "powerpoint")[0]["effects"] == ["weight 900 drawn as bold"])
small = lambda s: {"items": [{"type": "text", "x": 0, "y": 0, "w": 1, "h": 1, "runs": [{"t": "a", "family": "Inter", "spacing": s, "weight": 400}]}]}
check("even slight spacing is named, since it adds up along a line", fidelity.substitutions(small(0.42), {})[0]["effects"] == ["letter spacing not in Google Slides"])
check("no spacing, nothing named", fidelity.substitutions(small(0), {}) == [])
check("weights map to bold or regular", (fidelity.slides_weight(900), fidelity.slides_weight(600), fidelity.slides_weight(500)) == (700, 700, 400))
check("substitute font gets a Google Fonts link", "family=Roboto+Mono" in fidelity.substitute_links(m, brand))
check("no link for fonts Slides already has", fidelity.substitute_links({"items": [{"type": "text", "runs": [{"t": "a", "family": "Inter"}]}]}, {}) == "")

# ── font imports: a semicolon inside the URL must not cut the rule ─────────
print("font imports")
deck = '<style>@import url("https://fonts.googleapis.com/css2?family=Barlow:ital,wght@1,700;1,800&display=swap");\n@font-face{font-family:"X";src:url(x.woff2)}</style><link rel="stylesheet" href="assets/tokens.css">'
links, faces = fidelity.font_faces(deck)
check("whole @import kept", '1,700;1,800&display=swap");' in faces, faces)
check("@font-face kept", '@font-face{font-family:"X"' in faces)
check("stylesheet link kept", 'assets/tokens.css' in links)

# ── stage transform ─────────────────────────────────────────────────────────
print("stage transform")
html = '<html><head><title>t</title></head><body><deck-stage><section>a</section><section>b</section></deck-stage><script src="deck-stage.js"></script></body></html>'
t = stage.transform(html, "talks/x.html", 2, "measure")
check("runtime script removed", "deck-stage.js" not in t)
check("base href points at the deck's folder", '<base href="/talks/">' in t, t[:200])
check("slide 2 pinned", "section:not(:nth-of-type(2))" in t)
check("measure mode injects the measurer", "/__stage/measure.js" in t and 'id="__model"' in t)
check("effects mode adds the effects style", "__stage_effects" in stage.transform(html, "x.html", 1, "effects"))
fx = stage.transform(html, "x.html", 1, "effects")
check("effects mode restores decoration colours, after the effects styles", "__stage_fx_pseudo" in fx and fx.index("__stage_effects_late") < fx.index("__stage_fx_pseudo"))
check("raw and measure modes don't", "__stage_fx_pseudo" not in stage.transform(html, "x.html", 1, "raw") and "__stage_fx_pseudo" not in t)
check("raw mode adds neither", "__stage_effects" not in stage.transform(html, "x.html", 1, "raw") and "measure.js" not in stage.transform(html, "x.html", 1, "raw"))
check("root-level deck has base /", '<base href="/">' in stage.transform(html, "x.html", 1, "raw"))

# ── regions, explanations, verdicts ─────────────────────────────────────────
print("explain and verdict")
region = {"x": 100, "y": 100, "w": 200, "h": 100}
glow = {"label": "HEADLINE", "x": 110, "y": 110, "w": 180, "h": 80, "effects": ["glowing or shadowed text"]}
out = fidelity.explain([region], {"items": [], "unsupported": [glow]})
check("a region covered by a glow is explained", out[0]["reasons"] and out[0]["reasons"][0]["effect"] == "glowing or shadowed text", out)
big = {"x": 0, "y": 0, "w": 1920, "h": 1080}
out = fidelity.explain([big], {"items": [], "unsupported": [glow]})
check("a whole-slide region isn't explained by one small glow", out[0]["reasons"] == [], out)
out = fidelity.explain([region], {"items": [{"type": "text", "x": 120, "y": 120, "w": 50, "h": 20, "runs": [{"t": "near me"}]}], "unsupported": []})
check("unexplained region names the nearest text", out[0]["reasons"] == [] and out[0]["near"] == "near me", out)
check("no regions is faithful", fidelity.verdict(0.2, []) == "faithful")
check("tiny score, all explained, is faithful", fidelity.verdict(0.001, [{"reasons": [{"effect": "glowing or shadowed text"}]}]) == "faithful")
check("a small unexplained difference is still other (a missing bullet is a few pixels)", fidelity.verdict(0.0009, [{"reasons": []}]) == "other")
check("glow only is effects-only", fidelity.verdict(0.01, [{"reasons": [{"effect": "glowing or shadowed text"}]}]) == "effects-only")
check("CSS decoration is effects-only", fidelity.verdict(0.01, [{"reasons": [{"effect": 'CSS ::after decoration "×"'}]}]) == "effects-only")
check("unexplained is other", fidelity.verdict(0.01, [{"reasons": [{"effect": "glowing or shadowed text"}]}, {"reasons": []}]) == "other")
font_reason = {"effect": "font not in Google Slides: SF Mono → Roboto Mono"}
check("font substitution is approximated", fidelity.verdict(0.01, [{"reasons": [font_reason]}]) == "approximated")
check("glow plus a gradient approximation is approximated",
      fidelity.verdict(0.01, [{"reasons": [{"effect": "glowing or shadowed text"}]}, {"reasons": [{"effect": "gradient approximated: repeating gradient drawn once"}]}]) == "approximated")
check("an approximation that doesn't show is still named", fidelity.verdict(0.0, [], notes=["gradient approximated: radial spread"]) == "approximated")
check("an approximation doesn't excuse an unexplained region", fidelity.verdict(0.01, [{"reasons": [font_reason]}, {"reasons": []}], notes=["x"]) == "other")
check("treatments per verdict", fidelity.TREATMENTS["faithful"] == ["editable"] and "glow-behind" in fidelity.TREATMENTS["effects-only"]
      and "glow-behind" not in fidelity.TREATMENTS["other"] and "glow-behind" not in fidelity.TREATMENTS["approximated"])
check("approximated with a glow offers glow-behind", "glow-behind" in fidelity.treatments("approximated", [{"reasons": [{"effect": "glowing or shadowed text"}, font_reason]}]))
check("approximated without one doesn't", "glow-behind" not in fidelity.treatments("approximated", [{"reasons": [font_reason]}]))
rgn = {"x": 100, "y": 100, "w": 200, "h": 100}
box_approx = {"label": ".bar", "x": 100, "y": 100, "w": 200, "h": 100, "effects": ["gradient approximated: repeating gradient drawn once"], "visible": True, "background": False}
check("a visible approximation explains its region", fidelity.explain([rgn], {"items": [], "unsupported": [], "approximations": [box_approx]})[0]["reasons"])
check("an invisible one doesn't", not fidelity.explain([rgn], {"items": [], "unsupported": [], "approximations": [{**box_approx, "visible": False}]})[0]["reasons"])
bg_approx = {**box_approx, "x": 0, "y": 0, "w": 1920, "h": 1080, "background": True}
check("a background approximation explains bare background", fidelity.explain([rgn], {"items": [], "unsupported": [], "approximations": [bg_approx]})[0]["reasons"])
marker = {"label": "list marker •", "x": 1032, "y": 690, "w": 29, "h": 51, "effects": ["list marker position estimated"], "visible": True, "background": False, "quiet": True}
out = fidelity.explain([{"x": 1032, "y": 696, "w": 48, "h": 24}], {"items": [], "unsupported": [], "approximations": [marker]})
check("a bullet a pixel off is explained by its estimated position", out[0]["reasons"] and out[0]["reasons"][0]["effect"] == "list marker position estimated", out)
check("explained, so not other: faithful when tiny", fidelity.verdict(0.0001, out) == "faithful")
check("and approximated when it's big enough to see", fidelity.verdict(0.01, out) == "approximated")
check("a quiet approximation isn't a note on its own", fidelity.slide_notes({"approximations": [marker]}, {}) == [])
txt = {"type": "text", "x": 150, "y": 120, "w": 60, "h": 30, "runs": [{"t": "words"}]}
check("but not a region with text in it", not fidelity.explain([rgn], {"items": [txt], "unsupported": [], "approximations": [bg_approx]})[0]["reasons"])

# ── the image comparison ────────────────────────────────────────────────────
print("diff_regions")
try:
    from PIL import Image, ImageDraw
    tmp = Path(tempfile.mkdtemp(prefix="fid-test-"))
    a = Image.new("RGB", (480, 270), (10, 13, 27)); ImageDraw.Draw(a).rectangle([100, 100, 200, 160], fill=(255, 255, 255))
    a.save(tmp / "a.png"); a.save(tmp / "same.png")
    score, regions = fidelity.diff_regions(tmp / "a.png", tmp / "same.png")
    check("identical renders: no regions", regions == [] and score == 0, (score, regions))
    b = a.copy(); ImageDraw.Draw(b).rectangle([300, 40, 420, 120], fill=(236, 72, 153)); b.save(tmp / "b.png")
    score, regions = fidelity.diff_regions(tmp / "a.png", tmp / "b.png")
    check("an added shape is one region", len(regions) == 1, regions)
    r = regions[0] if regions else {"x": 0, "y": 0, "w": 0, "h": 0}
    check("the region covers the shape", r["x"] <= 300 and r["y"] <= 40 and r["x"] + r["w"] >= 420 and r["y"] + r["h"] >= 120, r)
    s1 = a.copy(); s1 = Image.new("RGB", (480, 270), (10, 13, 27)); ImageDraw.Draw(s1).rectangle([101, 100, 201, 160], fill=(255, 255, 255)); s1.save(tmp / "shift.png")
    score, regions = fidelity.diff_regions(tmp / "a.png", tmp / "shift.png")
    check("a one-pixel shift is not a loss", regions == [], regions)
except ImportError:
    print("  skip diff_regions (no Pillow)")

# ── export: which treatment a slide gets ────────────────────────────────────
print("export decide")
try:
    import export_pptx as ex
    glow_model = {"unsupported": [{"effects": ["glowing or shadowed text"]}]}
    svg_model = {"unsupported": [{"effects": ["svg element"]}]}
    check("--treat wins", ex.decide(1, {"1": "picture"}, {"1": "editable"}, None, glow_model, "ask") == ("picture", "--treat"))
    check("deck.json next", ex.decide(1, {}, {"1": "glow-behind"}, None, glow_model, "ask") == ("glow-behind", "deck.json"))
    check("faithful report is editable", ex.decide(1, {}, {}, {"verdict": "faithful"}, glow_model, "ask") == ("editable", "faithful"))
    check("ask leaves it undecided, editable", ex.decide(1, {}, {}, {"verdict": "effects-only"}, glow_model, "ask") == ("editable", "undecided (effects-only)"))
    check("auto: effects-only goes editable", ex.decide(1, {}, {}, {"verdict": "effects-only"}, glow_model, "auto")[0] == "editable")
    check("auto: other goes picture", ex.decide(1, {}, {}, {"verdict": "other"}, svg_model, "auto")[0] == "picture")
    check("no report: verdict from the model", ex.decide(1, {}, {}, None, svg_model, "auto")[0] == "picture")
    check("no report, no effects: faithful", ex.decide(1, {}, {}, None, {"unsupported": []}, "ask") == ("editable", "faithful"))
    check("auto: approximated goes editable", ex.decide(1, {}, {}, {"verdict": "approximated"}, {}, "auto") == ("editable", "auto (approximated)"))
    check("no report: approximations make it approximated",
          ex.decide(1, {}, {}, None, {"unsupported": [], "approximations": [{"effects": ["gradient approximated: x"]}]}, "ask") == ("editable", "undecided (approximated)"))
    try:
        from pptx import Presentation
        from pptx.oxml.ns import qn
        prs = Presentation(); sl = prs.slides.add_slide(prs.slide_layouts[6])
        ink, purple, pink = {"r": 10, "g": 13, "b": 27, "a": 1}, {"r": 100, "g": 44, "b": 169, "a": 0.5}, {"r": 255, "g": 41, "b": 143, "a": 1}
        ex.apply_gradient(sl.background.fill, {"kind": "radial", "shape": "circle", "at": {"x": 0.5, "y": 0.6},
                                               "stops": [{"c": ink, "p": 0.45}, {"c": purple, "p": 0.7}, {"c": pink, "p": 1}]})
        gf = sl.background.fill._xPr.find(qn("a:gradFill"))
        gs = gf.find(qn("a:gsLst")).findall(qn("a:gs"))
        check("export keeps all three stops", [x.get("pos") for x in gs] == ["45000", "70000", "100000"], [x.get("pos") for x in gs])
        check("export keeps stop colours", [x.find(qn("a:srgbClr")).get("val") for x in gs] == ["0A0D1B", "642CA9", "FF298F"])
        check("export keeps stop alpha", gs[1].find(qn("a:srgbClr")).find(qn("a:alpha")).get("val") == "50000")
        ftr = gf.find(qn("a:path")).find(qn("a:fillToRect"))
        check("radial centre at 50% 60%", (ftr.get("l"), ftr.get("t"), ftr.get("r"), ftr.get("b")) == ("50000", "60000", "50000", "40000"))
        check("path follows the stop list", gf.find(qn("a:gsLst")).getnext().tag == qn("a:path"))
        shp = sl.shapes.add_shape(1, 0, 0, 100, 100)
        ex.apply_gradient(shp.fill, {"kind": "linear", "angle": 180, "stops": [{"c": ink, "p": 0}, {"c": pink, "p": 1}]})
        lin = shp.fill._xPr.find(qn("a:gradFill")).find(qn("a:lin"))
        check("CSS 180deg (top to bottom) is DrawingML 90deg clockwise", lin.get("ang") == "5400000", lin.get("ang"))
        ex.apply_gradient(shp.fill, {"kind": "linear", "angle": 90, "stops": [{"c": ink, "p": 0}, {"c": pink, "p": 1}]})
        check("CSS 90deg (left to right) is DrawingML 0", shp.fill._xPr.find(qn("a:gradFill")).find(qn("a:lin")).get("ang") == "0")
        check("re-applying leaves one direction element", len(shp.fill._xPr.find(qn("a:gradFill")).findall(qn("a:lin"))) == 1)
        buf = __import__("io").BytesIO(); prs.save(buf)
        check("the file still saves", buf.tell() > 0)
    except ImportError:
        print("  skip apply_gradient (no python-pptx)")
    title = {"x": 346, "w": 1229, "align": "center", "wrap": False}
    check("a centred one-line title grows both ways, as far as its nearer edge allows", ex.text_span(title) == (1, 1919), ex.text_span(title))
    check("a left-aligned line grows right", ex.text_span({"x": 110, "w": 577, "align": "left", "wrap": False}) == (110, 1810))
    check("a right-aligned line grows left", ex.text_span({"x": 1500, "w": 300, "align": "right", "wrap": False}) == (0.0, 1800))
    check("wrapping text keeps its width", ex.text_span({"x": 110, "w": 577, "align": "left", "wrap": True}) == (110, 577))
    check("an off-centre title grows only as far as its nearer edge", ex.text_span({"x": 100, "w": 400, "align": "center", "wrap": False}) == (0, 600))
    try:
        from PIL import Image
        import io as _io
        def png(wh):
            b = _io.BytesIO(); Image.new("RGBA", wh, (236, 72, 153, 255)).save(b, "PNG"); return b.getvalue()
        big = ex.shrink(png((1024, 1024)), 150, 150)
        check("a 1024² icon in a 150px slot is embedded at twice the slot", Image.open(_io.BytesIO(big)).size == (300, 300), Image.open(_io.BytesIO(big)).size)
        small = png((200, 100))
        check("an image already small enough is left byte for byte", ex.shrink(small, 150, 150) == small)
        check("the shape is kept when shrinking", Image.open(_io.BytesIO(ex.shrink(png((2000, 1000)), 400, 400))).size == (800, 400))
        check("bytes that aren't an image pass through", ex.shrink(b"not an image", 10, 10) == b"not an image")
    except ImportError:
        print("  skip shrink (no Pillow)")
    check("speaker notes parse", ex.notes_of('<script type="application/json" id="speaker-notes">["one", "two"]</script>') == ["one", "two"])
    check("no notes is empty", ex.notes_of("<html></html>") == [])

    # one-version decks
    two = {"leave_behind": "talk.html", "mainstage": "talk-mainstage.html"}
    check("two versions: each file by its version", (fidelity.deck_file(two, "mainstage"), fidelity.deck_file(two, "leave-behind")) == ("talk-mainstage.html", "talk.html"))
    one = {"deck": "my-stack.html"}
    check('one version under "deck" answers for either', fidelity.deck_file(one, "mainstage") == "my-stack.html" == fidelity.deck_file(one, "leave-behind"))
    check("one version under a single key answers for the other too", fidelity.deck_file({"leave_behind": "x.html"}, "mainstage") == "x.html")
    check("two different files, neither for this version, is nothing", fidelity.deck_file({"deck": "a.html", "leave_behind": "b.html"}, "mainstage") is None)
    check("no file at all is nothing", fidelity.deck_file({"title": "t"}, "mainstage") is None)

    # pictures at twice the slide's size
    import common as common_mod
    seen, real_watch = {}, common_mod._watch
    common_mod._watch = lambda args, done, wait: (seen.setdefault("args", args), False)[1]
    common_mod.screenshot("chrome", "http://x", Path(tempfile.mkdtemp()) / "s.png", scale=2)
    common_mod._watch = real_watch
    check("screenshots can render at twice the slide's size", "--force-device-scale-factor=2" in seen.get("args", []), seen)
    check("the export renders at 2x", ex.RENDER_SCALE == 2)
    try:
        from PIL import Image
        big = Path(tempfile.mkdtemp()) / "raw2x.png"
        Image.new("RGBA", (3840, 2160), (0, 0, 0, 255)).save(big)
        crop = Image.open(ex.crop_png(big, 100, 50, 200, 100, scale=2))
        check("a cut-out from a 2x render covers the same box, at twice the pixels", crop.size == (400, 200), crop.size)
    except ImportError:
        print("  skip 2x crop (no Pillow)")
except SystemExit as e:
    print(f"  skip export decide ({e})")

print(f"\n{passed} passed, {failed} failed")
sys.exit(1 if failed else 0)

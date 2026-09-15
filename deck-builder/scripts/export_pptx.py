#!/usr/bin/env python3
"""Export a deck to a PowerPoint file that opens editable in Google Slides.

    python3 export_pptx.py talk/deck.json                         # the mainstage
    python3 export_pptx.py talk/deck.json --version leave-behind
    python3 export_pptx.py talk/deck.json --treat 1=picture,17=glow-behind --save
    python3 export_pptx.py talk/deck.json --out /tmp/check.pptx

Each slide gets one of three treatments:

  editable      native shapes from the measured slide: text boxes with their font, size,
                weight, colour and alignment; filled, bordered and rounded boxes; gradients;
                images. What an editable copy can't draw (an SVG, a canvas) is placed as a
                picture cut from the render, so nothing goes missing.
  glow-behind   the same, with the effects the shapes can't carry (glows, shadows, CSS
                decorations) rendered on a transparent picture underneath. The text stays
                editable; the glow keeps showing the old words if you change them.
  picture       the slide as rendered, exact, not editable.

Where the treatment comes from, first match wins: --treat, then deck.json › slides, then the
fidelity report if one exists (faithful slides are editable), then pptx.unfaithful: "auto"
sends effects-only and approximated slides editable and anything else to picture; "ask"
exports editable and lists the slides still waiting on a decision.

Gradients keep every stop, its alpha, and a radial's centre. --save writes --treat choices into deck.json.

Speaker notes come from the deck's <script id="speaker-notes">. Needs Chrome and python-pptx.
"""
import argparse, base64, io, json, re, sys, urllib.request
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import common
from stage import Stage
import fidelity

EMU_PER_PX = 6350       # 1920 px across a 13.333-inch slide
PT_PER_PX = 0.5         # 1920 px across 960 pt
# Slides rendered as pictures (the picture treatment, the effects behind glow-behind text, cut-outs
# of what shapes can't draw) are rendered at twice the slide's size. At 1× a picture slide goes
# soft on a projector or a 4K screen.
RENDER_SCALE = 2

def need_pptx():
    try:
        import pptx  # noqa: F401
    except ImportError:
        sys.exit("export_pptx: needs python-pptx (pip install python-pptx)")

def notes_of(deck_html):
    m = re.search(r'<script[^>]*id="speaker-notes"[^>]*>(.*?)</script>', deck_html, re.S)
    try: return json.loads(m.group(1)) if m else []
    except json.JSONDecodeError: return []

def px(v): return int(round(v * EMU_PER_PX))

def set_alpha(color_format_parent_xml, alpha):
    """Add <a:alpha> to the srgbClr under a fill or run properties element."""
    if alpha is None or alpha >= 0.999: return
    from pptx.oxml.ns import qn
    for clr in color_format_parent_xml.iter(qn("a:srgbClr")):
        for old in clr.findall(qn("a:alpha")): clr.remove(old)
        a = clr.makeelement(qn("a:alpha"), {"val": str(int(round(alpha * 100000)))})
        clr.append(a)

def rgb(c):
    from pptx.dml.color import RGBColor
    return RGBColor(int(c["r"]), int(c["g"]), int(c["b"]))

def apply_gradient(fill, g):
    """Every stop of the CSS gradient, with its alpha, written straight into DrawingML.
    python-pptx's gradient_stops only exposes the two stops it starts with, and its
    gradient_angle counts counter-clockwise while <a:lin ang> is clockwise."""
    from pptx.oxml.ns import qn
    fill.gradient()
    gradFill = fill._xPr.find(qn("a:gradFill"))
    gsLst = gradFill.find(qn("a:gsLst"))
    for gs in list(gsLst): gsLst.remove(gs)
    for s in g["stops"]:
        gs = gsLst.makeelement(qn("a:gs"), {"pos": str(int(round(max(0.0, min(1.0, s["p"])) * 100000)))})
        c = s["c"]
        clr = gs.makeelement(qn("a:srgbClr"), {"val": "%02X%02X%02X" % (int(c["r"]), int(c["g"]), int(c["b"]))})
        if c.get("a", 1) < 0.999:
            clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(round(c["a"] * 100000)))}))
        gs.append(clr); gsLst.append(gs)
    for old in gradFill.findall(qn("a:lin")) + gradFill.findall(qn("a:path")): gradFill.remove(old)
    if g["kind"] == "linear":
        # CSS 90deg runs left→right, DrawingML 0; both then turn clockwise
        el = gradFill.makeelement(qn("a:lin"), {"ang": str(int(round(((g["angle"] - 90) % 360) * 60000))), "scaled": "0"})
    else:
        at = g.get("at") or {"x": 0.5, "y": 0.5}
        el = gradFill.makeelement(qn("a:path"), {"path": "circle"})
        el.append(el.makeelement(qn("a:fillToRect"), {"l": str(int(round(at["x"] * 100000))), "t": str(int(round(at["y"] * 100000))),
                                                     "r": str(int(round((1 - at["x"]) * 100000))), "b": str(int(round((1 - at["y"]) * 100000)))}))
    gsLst.addnext(el)

def fetch(url):
    if url.startswith("data:"):
        return base64.b64decode(url.split(",", 1)[1])
    with urllib.request.urlopen(url, timeout=30) as r: return r.read()

def crop_png(png_path, x, y, w, h, scale=1):
    """A cut-out of a render. x, y, w, h are slide pixels; `scale` is how many render pixels
    make one slide pixel, so a 2x render gives the same area at twice the detail."""
    from PIL import Image
    im = Image.open(png_path).convert("RGBA")
    x, y, w, h = x * scale, y * scale, w * scale, h * scale
    box = (max(0, int(x)), max(0, int(y)), min(im.width, int(x + w)), min(im.height, int(y + h)))
    buf = io.BytesIO(); im.crop(box).save(buf, "PNG"); buf.seek(0); return buf

def shrink(data, placed_w, placed_h, over=2.0):
    """Image bytes no larger than `over`× their placed size, shape kept. Generated icons arrive
    at 1024² and sit in a 150px slot; embedding them whole took a 17-slide deck from 5 MB to
    21 MB. Twice the placed size stays sharp on a 4K screen. Anything that isn't an image
    Pillow can read, or is already small enough, passes through unchanged."""
    try:
        from PIL import Image
        im = Image.open(io.BytesIO(data))
        tw, th = max(1, round(placed_w * over)), max(1, round(placed_h * over))
        if im.width <= tw and im.height <= th: return data
        fmt = "JPEG" if im.format == "JPEG" else "PNG"
        im = im.convert("RGB" if fmt == "JPEG" else "RGBA")
        im.thumbnail((tw, th), Image.LANCZOS)
        buf = io.BytesIO()
        im.save(buf, fmt, optimize=True, **({"quality": 90} if fmt == "JPEG" else {}))
        return buf.getvalue()
    except Exception:
        return data

def add_image(slide, it):
    data = fetch(it["src"])
    try:
        from PIL import Image
        im = Image.open(io.BytesIO(data)); nw, nh = im.size
        if im.format not in ("PNG", "JPEG", "GIF", "BMP", "TIFF"):
            buf = io.BytesIO(); im.convert("RGBA").save(buf, "PNG"); data = buf.getvalue()
    except Exception:
        nw, nh = it.get("nw") or 1, it.get("nh") or 1
    x, y, w, h = it["x"], it["y"], it["w"], it["h"]
    fit = it.get("fit") or "fill"
    if fit in ("contain", "scale-down") and nw and nh:
        k = min(w / nw, h / nh); dw, dh = nw * k, nh * k
        pic = slide.shapes.add_picture(io.BytesIO(shrink(data, dw, dh)), px(x + (w - dw) / 2), px(y + (h - dh) / 2), px(dw), px(dh))
    elif fit == "cover" and nw and nh:
        k = max(w / nw, h / nh); dw, dh = nw * k, nh * k
        pic = slide.shapes.add_picture(io.BytesIO(shrink(data, dw, dh)), px(x), px(y), px(w), px(h))
        cx, cy = (dw - w) / 2 / dw, (dh - h) / 2 / dh
        pic.crop_left = pic.crop_right = cx; pic.crop_top = pic.crop_bottom = cy
    else:
        pic = slide.shapes.add_picture(io.BytesIO(shrink(data, w, h)), px(x), px(y), px(w), px(h))
    pic.name = "Image — " + it["src"].rsplit("/", 1)[-1][:40]
    return pic

def add_box(slide, it):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Pt
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if it.get("radius") else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, px(it["x"]), px(it["y"]), px(it["w"]), px(it["h"]))
    if it.get("radius"):
        s.adjustments[0] = max(0.0, min(0.5, it["radius"] / max(1.0, min(it["w"], it["h"]))))
    if it.get("gradient"):
        apply_gradient(s.fill, it["gradient"])
    elif it.get("fill"):
        s.fill.solid(); s.fill.fore_color.rgb = rgb(it["fill"]); set_alpha(s.fill._xPr, it["fill"]["a"] * it.get("opacity", 1))
    else:
        s.fill.background()
    if it.get("border"):
        s.line.color.rgb = rgb(it["border"]["c"]); s.line.width = Pt(it["border"]["w"] * PT_PER_PX)
        set_alpha(s.line._ln, it["border"]["c"]["a"])
        if it["border"]["dash"]: s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    s.name = "Box"
    return s

def text_span(it, slide_w=common.W):
    """(x, width) for a text box. Text that doesn't wrap in the browser gets room to spare:
    Google Slides wraps even a wrap="none" box at its width on import, and sets text wider than
    Chrome does (no letter spacing, bold for black, its own metrics), so a box cut to the glyphs
    rewraps. The box grows away from its alignment edge, as far as the slide allows."""
    x, w = it["x"], max(it["w"], 1)
    if it.get("wrap"): return x, w
    align = it.get("align")
    if align == "center":
        room = max(0.0, min(x, slide_w - (x + w)))
        return x - room, w + 2 * room
    if align == "right":
        return 0.0, max(w, x + w)
    return x, max(w, slide_w - x)

def add_text(slide, it, brand):
    from pptx.util import Pt
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.oxml.ns import qn
    bx, bw = text_span(it)
    tb = slide.shapes.add_textbox(px(bx), px(it["y"]), px(bw), px(max(it["h"], 1)))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = bool(it.get("wrap"))
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = {"middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}.get(it.get("valign"), MSO_ANCHOR.TOP)
    align = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}.get(it.get("align"), PP_ALIGN.LEFT)
    paras = [[]]
    for r in it["runs"]:
        if r.get("br"): paras.append([])
        else: paras[-1].append(r)
    first_text = ""
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        # a multiple of the font's single spacing: Google Slides drops exact point spacing on import
        if it.get("lineHeight") and it.get("normalLH"): p.line_spacing = round(it["lineHeight"] / it["normalLH"], 3)
        elif it.get("lineHeight"): p.line_spacing = Pt(it["lineHeight"] * PT_PER_PX)
        p.space_before = p.space_after = Pt(0)
        for r in runs:
            run = p.add_run(); run.text = r["t"]; first_text += r["t"]
            f = run.font
            f.name = fidelity.font_for(r["family"], brand)
            f.size = Pt(max(1, r["size"] * PT_PER_PX))
            f.bold = r["weight"] >= 600
            f.italic = bool(r["italic"])
            if r.get("underline"): f.underline = True
            if r.get("color"):
                f.color.rgb = rgb(r["color"])
                set_alpha(run._r.get_or_add_rPr(), r["color"]["a"] * it.get("opacity", 1))
            if r.get("spacing"):
                run._r.get_or_add_rPr().set("spc", str(int(round(r["spacing"] * PT_PER_PX * 100))))
    tb.name = "Text — " + first_text.strip()[:40]
    return tb

def add_notes(slide, text):
    if text: slide.notes_slide.notes_text_frame.text = text

def decide(n, cli, saved, report_row, model, cfg_mode):
    """(treatment, where it came from)"""
    if str(n) in cli: return cli[str(n)], "--treat"
    if str(n) in saved: return saved[str(n)], "deck.json"
    if report_row:
        v = report_row["verdict"]
    else:
        kinds = {e for u in (model or {}).get("unsupported", []) for e in u["effects"]}
        approx = (model or {}).get("approximations") or []
        v = ("other" if kinds and not all(k.startswith(fidelity.EFFECT_KINDS) for k in kinds)
             else "approximated" if approx else "effects-only" if kinds else "faithful")
    if v == "faithful": return "editable", "faithful"
    if cfg_mode == "auto": return ("editable" if v in ("effects-only", "approximated") else "picture"), f"auto ({v})"
    return "editable", f"undecided ({v})"

def export(deck_json, version="mainstage", out=None, treat=None, save=False):
    need_pptx()
    from pptx import Presentation
    from pptx.util import Emu
    folder, deck, cfg = fidelity.load_deck(deck_json, version)
    file_rel = fidelity.deck_file(deck, version)
    if not file_rel: sys.exit(f'export_pptx: deck.json names no file for the {version} (a one-version deck puts its file under "deck")')
    deck_html = (folder / file_rel).read_text()
    brand = common.load_brand(folder)
    chrome = common.chrome_path(brand)
    cfg_mode = ((cfg.get("pptx") or {}).get("unfaithful")) or "ask"
    notes = notes_of(deck_html)
    work = folder / ".deck-builder" / "export" / version; work.mkdir(parents=True, exist_ok=True)
    report_path = folder / ".deck-builder" / "fidelity" / version / "report.json"
    report = {r["slide"]: r for r in json.loads(report_path.read_text())["slides"]} if report_path.is_file() else {}
    cli = dict(kv.split("=", 1) for kv in treat.split(",")) if treat else {}
    for v in cli.values():
        if v not in ("editable", "glow-behind", "picture"): sys.exit(f"export_pptx: unknown treatment {v!r}")
    import decks_config
    saved = decks_config.slides_for(deck, version)

    prs = Presentation(); prs.slide_width, prs.slide_height = Emu(px(common.W)), Emu(px(common.H))
    blank = prs.slide_layouts[6]
    n_slides = fidelity.slide_count(deck_html)
    log = []
    with Stage(folder) as st:
        for n in range(1, n_slides + 1):
            # one render at twice the slide's size serves the picture treatment and the cut-outs
            raw = common.screenshot(chrome, st.url(file_rel, n, "raw"), work / f"{n:02d}-raw.png", scale=RENDER_SCALE)
            model = fidelity.measure(st, chrome, file_rel, n)
            treatment, why = decide(n, cli, saved, report.get(n), model, cfg_mode)
            if (not model or "error" in model) and treatment != "picture":
                treatment, why = "picture", "could not measure"
            slide = prs.slides.add_slide(blank)
            if treatment == "picture":
                if not raw: sys.exit(f"export_pptx: slide {n} did not render")
                pic = slide.shapes.add_picture(str(raw), 0, 0, prs.slide_width, prs.slide_height); pic.name = f"Slide {n} — picture"
            else:
                bg = model.get("background") or {}
                if bg.get("gradient"): apply_gradient(slide.background.fill, bg["gradient"])
                elif bg.get("fill"): slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(bg["fill"])
                if treatment == "glow-behind":
                    fx = common.screenshot(chrome, st.url(file_rel, n, "effects"), work / f"{n:02d}-effects.png", transparent=True, scale=RENDER_SCALE)
                    if fx:
                        p = slide.shapes.add_picture(str(fx), 0, 0, prs.slide_width, prs.slide_height); p.name = "Effects — glows and shadows (picture)"
                for it in model["items"]:
                    try:
                        if it["type"] == "box": add_box(slide, it)
                        elif it["type"] == "text": add_text(slide, it, brand)
                        elif it["type"] == "image": add_image(slide, it)
                        elif it["type"] == "raster" and raw:
                            p = slide.shapes.add_picture(crop_png(raw, it["x"], it["y"], it["w"], it["h"], scale=RENDER_SCALE), px(it["x"]), px(it["y"]), px(it["w"]), px(it["h"]))
                            p.name = f"Picture — {it.get('why', 'element')} the editable copy can't draw"
                    except Exception as ex:
                        log.append(f"slide {n}: skipped a {it['type']} ({ex})")
            add_notes(slide, notes[n - 1] if n - 1 < len(notes) else "")
            log.append((n, treatment, why))
    out = Path(out) if out else folder / (Path(file_rel).stem + ".pptx")
    prs.save(out)
    if save and cli:
        dj = Path(deck_json).resolve(); d = json.loads(dj.read_text())
        s = d.setdefault("slides", {})
        # a deck that keeps choices per version gets this version's map updated, not a flat one
        (s.setdefault(version, {}) if any(isinstance(v, dict) for v in s.values()) else s).update(cli)
        dj.write_text(json.dumps(d, indent=2, ensure_ascii=False) + "\n")
    return out, log

if __name__ == "__main__":
    a = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    a.add_argument("deck"); a.add_argument("--version", choices=["mainstage", "leave-behind"], default="mainstage")
    a.add_argument("--out"); a.add_argument("--treat", help="e.g. 1=picture,17=glow-behind")
    a.add_argument("--save", action="store_true", help="write --treat choices into deck.json")
    args = a.parse_args()
    out, log = export(args.deck, args.version, args.out, args.treat, args.save)
    counts = {}
    undecided = []
    for row in log:
        if isinstance(row, str): print("  " + row); continue
        n, t, why = row; counts[t] = counts.get(t, 0) + 1
        if why.startswith("undecided"): undecided.append(n)
        print(f"  slide {n:>2}  {t:11}  {why}")
    print(f"{out}  ({sum(counts.values())} slides: " + ", ".join(f"{v} {k}" for k, v in counts.items()) + ")")
    if undecided:
        print(f"waiting on a decision: slides {', '.join(map(str, undecided))} went editable by default; "
              f"choose with --treat N=editable|glow-behind|picture --save")
